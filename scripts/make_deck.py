"""Builds docs/Krishi_Sahay_ideation_deck.pptx (12 slides, 16:9) from the project's own documents.

Every number comes from docs/solution_writeup.md, reports/impact_report.md or docs/deployment_plan.md.
Team names are placeholders: edit them on the last slide.   Run: py -3.11 scripts/make_deck.py
"""

import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "docs", "Krishi_Sahay_ideation_deck.pptx")
DIAGRAM = os.path.join(ROOT, "docs", "architecture_ppt.png")

DARK = RGBColor(0x0B, 0x3D, 0x2E)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
LIGHT = RGBColor(0xEE, 0xF7, 0xF0)
INK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xD9, 0x77, 0x06)
TEAL = RGBColor(0x08, 0x91, 0xB2)
BLUE = RGBColor(0x25, 0x63, 0xEB)
RED = RGBColor(0xB9, 0x1C, 0x1C)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 12


def rect(slide, x, y, w, h, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: list of (text, size, bold, color) or (text, size, bold, color, space_after)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    for i, p in enumerate(paras):
        t, size, bold, color = p[:4]
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(p[4] if len(p) > 4 else 4)
        r = para.add_run()
        r.text = t
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return tb


def bullets(items, size=16, color=INK, gap=6):
    return [("•  " + i, size, False, color, gap) for i in items]


def base(n, title, kicker=None):
    s = prs.slides.add_slide(BLANK)
    bg = s.background.fill
    bg.solid()
    bg.fore_color.rgb = WHITE
    rect(s, 0, 0, 13.333, 0.12, DARK, shape=MSO_SHAPE.RECTANGLE)
    if kicker:
        text(s, 0.55, 0.28, 9, 0.35, [(kicker.upper(), 12, True, GREEN)])
    text(s, 0.55, 0.58, 12.2, 0.9, [(title, 26, True, DARK)])
    text(s, 0.55, 7.08, 9, 0.3, [("Krishi Sahay  |  Sustainable Agriculture: Energy, Water & Productivity  |  Ideation round", 10, False, GREY)])
    text(s, 12.0, 7.08, 0.8, 0.3, [(f"{n} / {TOTAL}", 10, False, GREY)], align=PP_ALIGN.RIGHT)
    return s


def card(s, x, y, w, h, head, body, accent=GREEN, size=14, fill=LIGHT):
    size = size + 2 if h >= 3.0 else size
    rect(s, x, y, w, h, fill)
    rect(s, x, y + 0.12, 0.07, h - 0.24, accent, shape=MSO_SHAPE.RECTANGLE)
    text(s, x + 0.2, y + 0.1, w - 0.3, 0.45, [(head, 17, True, DARK)])
    text(s, x + 0.2, y + 0.58, w - 0.3, h - 0.65, bullets(body, size, INK, 5))


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# 1 ---------------------------------------------------------------- title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = DARK
rect(s, 0.6, 1.55, 0.12, 2.2, GREEN, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.95, 1.4, 11.5, 1.2, [("Krishi Sahay", 60, True, WHITE)])
text(s, 0.95, 2.55, 11.5, 1.3, [("Plan the crop to the water, schedule the water and power, sell what you grow.", 26, False, RGBColor(0xD1, 0xFA, 0xE5))])
text(s, 0.95, 4.2, 11.5, 1.2, [
    ("Challenge: Sustainable Agriculture: Energy, Water & Productivity", 18, True, WHITE),
    ("Ideation round submission  |  A voice-first decision platform for Indian smallholders, in their own language", 16, False, RGBColor(0xA7, 0xF3, 0xD0))])
text(s, 0.95, 6.3, 11.5, 0.5, [("Team: [Team name]", 14, False, RGBColor(0xA7, 0xF3, 0xD0))])
notes(s, "Open with one sentence: we connect the crop, its water and power, the climate it will actually face, and the buyer, for one farmer's own land.")

# 2 ---------------------------------------------------------------- problem
s = base(2, "Farm power and water are large and unmanaged, and crop choice ignores both", "Problem understanding")
card(s, 0.55, 1.65, 4.0, 3.7, "Water and energy", [
    "Farm use is roughly 19–23% of India's electricity",
    "Groundwater supplies about 48% of irrigation water",
    "About 22 million electric pumps (2021)",
    "Irrigation timing is mostly habit: fixed-date flood watering ignores rain and soil"], TEAL, 14)
card(s, 4.67, 1.65, 4.0, 3.7, "Crop and timing", [
    "Recommenders rank by soil match or this year's rain: a dry spell or flood later in the season is invisible",
    "Costs and prices swing, so a 'best crop' can lose money in a bad year",
    "Power often comes at night; daytime supply is a stated promise, not a given"], AMBER, 14)
card(s, 8.79, 1.65, 4.0, 3.7, "Waste and reach", [
    "Post-harvest losses are 15–20% of produce value (challenge brief)",
    "Farmers plant what paid last year; everyone does, the glut is wasted",
    "Advice that needs English typing misses most smallholders; chatbots answer questions but do not build a plan from the farmer's land, water and buyers"], RED, 14)
rect(s, 0.55, 5.55, 12.24, 1.25, DARK)
text(s, 0.8, 5.65, 11.8, 1.1, [
    ("Nobody connects the crop, its water and power, the climate it will face, and the buyer, for one farmer's own land.", 20, True, WHITE),
    ("Shares are from secondary compilations and differ by source; sources are listed in the write-up.", 12, False, RGBColor(0xA7, 0xF3, 0xD0))], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Statistics are secondary compilations; the write-up lists the sources. Do not overstate them.")

# 3 ---------------------------------------------------------------- journey + screens
s = base(3, "User journey: from the farmer's location to a committed buyer, in their language", "Proposed solution, user journey and working prototype")
steps = [("1", "Location", "GPS finds the farm and names the village; or search a place or type coordinates"),
         ("2", "Water, land, soil", "How long the well lasts, acres, and soil from a test or a satellite estimate"),
         ("3", "Sowing date", "The farmer picks the date; all weather and risk numbers use it"),
         ("4", "Crops", "Best crops ranked by profit in a usual year and in a bad year"),
         ("5", "Plan and My farm", "Acres per crop, plus a watering schedule with energy use"),
         ("6", "Sell", "Buyers within 50 km with prices; commit a quantity to one"),
         ("7", "Ask by voice", "Ask in your own language and hear the answer as audio")]
w = 1.66
for i, (n, h, b_) in enumerate(steps):
    x = 0.55 + i * (w + 0.105)
    rect(s, x, 1.4, w, 1.65, LIGHT)
    rect(s, x + 0.08, 1.48, 0.34, 0.34, GREEN, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.08, 1.49, 0.34, 0.32, [(n, 13, True, WHITE)], align=PP_ALIGN.CENTER)
    text(s, x + 0.45, 1.46, w - 0.48, 0.4, [(h, 12.5, True, DARK)])
    text(s, x + 0.08, 1.9, w - 0.14, 1.15, [(b_, 11.5, False, INK)])
text(s, 0.55, 3.12, 12.2, 0.3, [("Screens from the working prototype (demo farmer near Kandi, Sangareddy; demo buyers and prices)", 13, True, GREEN)])
caps = [("crops", "Crop list for the farm: profit in a usual and a bad year, risk level, and buyers nearby"),
        ("plan", "How many acres of each crop to grow, with the profit range for the whole plan"),
        ("schedule", "Watering schedule: dates, water depth, volume, pump hours and electricity use"),
        ("sell", "Sell screen: every buyer within 50 km, with crop, quantity, price and dates")]
iw = 2.5
ih = iw * 775 / 635
for i, (k, cap) in enumerate(caps):
    x = 0.55 + i * 3.25
    pic = s.shapes.add_picture(os.path.join(ROOT, "docs", "screens", k + ".png"), Inches(x), Inches(3.45), width=Inches(iw))
    pic.line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    pic.line.width = Pt(1)
    text(s, x, 3.45 + ih + 0.04, 2.9, 0.6, [(cap, 11.5, False, INK)])
notes(s, "Screens are from the working prototype (demo farmer, Kandi, Sangareddy; buyers and prices are synthetic demo data).")

# 4 ---------------------------------------------------------------- alignment
s = base(4, "How it answers the challenge: water, energy, productivity and post-harvest loss", "Alignment to the selected challenge")
cols = [("Water", TEAL, [("What the farmer gets", "A day-by-day watering schedule from the sowing date, adjusted to rain and the 15-day forecast."),
                         ("How it is worked out", "A soil-water balance for each crop, using 30 years of weather at the farm (FAO-56 method)."),
                         ("How we measure it", "Cubic metres of water per acre, compared with flood irrigation on fixed dates.")]),
        ("Energy", AMBER, [("What the farmer gets", "Pump hours, electricity (kWh) and CO₂ for each crop's watering plan, and the right pump size."),
                           ("How it is worked out", "Water volume × pump lift ÷ pump efficiency gives kWh; kWh × the grid factor gives CO₂."),
                           ("How we measure it", "kWh and kg of CO₂ per acre per season.")]),
        ("Productivity", GREEN, [("What the farmer gets", "Crops ranked by profit, including a bad year, so dry-season failures are avoided."),
                                 ("How it is worked out", "Each crop is grown virtually in the last 30 seasons at that farm, on that land's own water supply."),
                                 ("How we measure it", "Relative yield and rupees of profit per acre.")]),
        ("Post-harvest loss", RED, [("What the farmer gets", "Wholesalers' priced needs within 50 km before sowing, and a warning if neighbours plan the same crop. Less crop waste."),
                                    ("How it is worked out", "Matching by distance (Haversine) and a planting registry that adds up nearby plans."),
                                    ("How we measure it", "Share of the harvest with a confirmed buyer (to be measured in the pilot).")])]
for i, (h, c, secs) in enumerate(cols):
    x = 0.55 + i * 3.1
    rect(s, x, 1.55, 2.95, 0.55, c)
    text(s, x, 1.58, 2.95, 0.5, [(h, 19, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    rect(s, x, 2.1, 2.95, 4.0, LIGHT)
    paras = []
    for sh, body_ in secs:
        paras.append((sh.upper(), 11, True, c, 1))
        paras.append((body_, 14, False, INK, 9))
    text(s, x + 0.1, 2.18, 2.75, 3.9, paras)
rect(s, 0.55, 6.25, 12.24, 0.65, DARK)
text(s, 0.75, 6.27, 11.9, 0.6, [("Water → energy → money: the water a crop needs is converted into pump hours, kWh, CO₂ and rupees for that farmer.", 15, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Each column says what the farmer gets, how it is computed and how it is measured.")

# 5 ---------------------------------------------------------------- architecture
s = base(5, "Architecture: how data, water, energy and money flow through four layers", "Architecture and design")
if os.path.exists(DIAGRAM):
    s.shapes.add_picture(DIAGRAM, Inches(0.55), Inches(1.45), width=Inches(10.15))
text(s, 10.85, 1.5, 2.05, 5.4, [("How to read it", 15, True, DARK),
                                 ("Rows run left to right: user, app screen, engine, outside data.", 12, False, INK, 8),
                                 ("Blue: data. Teal: water. Amber: energy and CO₂. Green: money and price.", 12, False, INK, 8),
                                 ("Today the stores are JSON files; production moves each to a database table.", 12, False, INK, 8),
                                 ("No hardware is needed: soil moisture is modelled; a sensor can plug in later.", 12, False, INK)])
notes(s, "The full-size diagram is docs/architecture_ppt.png (3840x2160). Replace this picture with it if you want more detail.")

# 6 ---------------------------------------------------------------- technical approach
s = base(6, "Technical approach: a per-farm simulation, not a district average", "Architecture and design")
card(s, 0.55, 1.6, 6.05, 1.75, "Climate risk and irrigation", [
    "Daily FAO-56 soil-water balance for every crop, for each of the last 30 seasons at that farm (ERA5 weather)",
    "Yield loss from FAO-33; result is a median, a bad year (10th percentile) and the share of poor years"], TEAL, 13)
card(s, 6.75, 1.6, 6.05, 1.75, "Profit and risk", [
    "Profit = yield × farm-gate price × (1 − loss) − cost, from a cited table (CACP, DES/NHB, Agmarknet, PIB MSP, NABCONS)",
    "Estimated numbers are flagged, discounted and never ranked first"], GREEN, 13)
card(s, 0.55, 3.5, 6.05, 1.75, "Water to power to money", [
    "E = ρ g H V / (3.6·10⁶ η) gives kWh; CO₂ uses the CEA grid factor 0.727 kg/kWh (diesel 2.68 kg/L)",
    "Pump size and payback use the farmer's own tariff and subsidy inputs"], AMBER, 13)
card(s, 6.75, 3.5, 6.05, 1.75, "Voice and matching", [
    "Speech to text, translate, Llama-3.3 answer with the farmer's land and plan as context, translate back, audio",
    "Buyers within 50 km by Haversine; priced offers; two-step confirmation with a reference code"], BLUE, 13)
rect(s, 0.55, 5.45, 12.24, 1.45, WHITE, GREY)
text(s, 0.75, 5.5, 11.9, 1.4, [("Open data in, no paid APIs needed to start", 15, True, DARK),
                              ("Open-Meteo (ERA5 + forecast)  |  ISRIC SoilGrids (soil pH, organic carbon, texture)  |  Agmarknet via CEDA (mandi prices)  |  OpenStreetMap (place names)  |  Together AI (Whisper, Llama-3.3)  |  Google Translate + gTTS (pilot only)", 13, False, INK, 6),
                              ("Stated assumptions: soil water-holding by class, 60 mm per flood irrigation, 40 m pump lift at 35% efficiency. Seasonal monsoon forecasts are deliberately not used to rank crops.", 12, False, GREY)])
notes(s, "Say plainly which numbers are assumptions. The honest framing is a strength with judges.")

# 7 ---------------------------------------------------------------- innovation
s = base(7, "Innovation: it joins what existing tools keep apart", "Idea quality and originality")
rows = [(" ", "Typical tools", "Krishi Sahay"),
        ("Crop advice", "Soil match or this year's rain", "30 seasons of weather at the exact farm, from the farmer's own sowing date"),
        ("Money", "Average price", "Usual-year and bad-year profit, price floor (MSP), glut caps"),
        ("Water and power", "Separate or absent", "Water to kWh to CO₂ to rupees, and a village roll-up"),
        ("Selling", "Mandi price apps", "Priced buyer offers within 50 km before sowing"),
        ("Access", "English text chatbots", "Voice in the farmer's language, history remembered"),
        ("Trust", "Confident single answer", "Ranges, plain reasons, estimates flagged, 'ask your KVK'")]
tbl = s.shapes.add_table(len(rows), 3, Inches(0.55), Inches(1.6), Inches(12.24), Inches(4.5)).table
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(2.1), Inches(3.9), Inches(6.24)
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.runs[0].font.size = Pt(15 if r else 16)
        p.runs[0].font.bold = (r == 0 or c == 0)
        p.runs[0].font.name = "Calibri"
        p.runs[0].font.color.rgb = WHITE if r == 0 else (DARK if c == 0 else INK)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK if r == 0 else (LIGHT if r % 2 else WHITE)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
rect(s, 0.55, 6.3, 12.24, 0.6, GREEN)
text(s, 0.75, 6.32, 11.9, 0.55, [("The loop that is new: what to grow depends on what nearby buyers want and what neighbours plan to plant.", 16, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "'Typical tools' is a generalisation drawn from the tools reviewed in the write-up (crop recommenders, Kisan e-Mitra-style chatbots, mandi apps). Soften it if challenged on a specific product.")

# 8 ---------------------------------------------------------------- impact
s = base(8, "Impact: up to 57% less water per acre at equal yield, vs a stated baseline", "Impact and measurability")
cd = CategoryChartData()
cd.categories = ["Cotton", "Tomato", "Onion"]
cd.add_series("Scheduling only", (17, 35, -20))
cd.add_series("Scheduling + drip", (44, 57, 20))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.55), Inches(1.6), Inches(6.6), Inches(4.0), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(13)
ch.has_title = True
ch.chart_title.text_frame.text = "Water saved per acre vs baseline (%)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(15)
ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
ch.category_axis.tick_labels.font.size = Pt(14)
ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
ch.value_axis.tick_labels.font.size = Pt(12)
ch.value_axis.has_major_gridlines = False
for ser, col in zip(ch.plots[0].series, (TEAL, GREEN)):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = col
    ser.invert_if_negative = False
ch.plots[0].has_data_labels = True
ch.plots[0].data_labels.font.size = Pt(13)
ch.plots[0].data_labels.font.bold = True
ch.plots[0].data_labels.number_format = '0"%"'
ch.plots[0].data_labels.number_format_is_linked = False
ch.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
text(s, 0.55, 5.65, 6.6, 0.9, [("Negative = uses more water. Onion with scheduling alone uses 20% more water but gains +28% relative yield.", 12, False, GREY)])
rect(s, 7.35, 1.6, 5.45, 1.55, LIGHT)
text(s, 7.5, 1.65, 5.2, 1.5, [("Baseline (stated)", 15, True, DARK),
                              ("The crop's usual number of flood irrigations on fixed dates, 60 mm each, no rain adjustment; 40 m lift, 35% pump efficiency. Yavatmal, black soil, 30 seasons of ERA5 weather.", 13, False, INK)])
rect(s, 7.35, 3.25, 5.45, 1.55, LIGHT)
text(s, 7.5, 3.3, 5.2, 1.5, [("Productivity gains, not savings", 15, True, DARK),
                             ("Where usual practice under-waters, the schedule spends more water for yield: wheat +7%, tur +17%, mustard +51%, chana +59% (relative yield).", 13, False, INK)])
rect(s, 7.35, 4.9, 5.45, 1.55, LIGHT)
text(s, 7.5, 4.95, 5.2, 1.5, [("Energy and waste", 15, True, DARK),
                              ("Cotton: about 202 kWh and 147 kg CO₂ saved per acre per season. Post-harvest loss at stake: tomato about ₹24,700 per acre (upper bound).", 13, False, INK)])
rect(s, 0.55, 6.5, 12.24, 0.5, DARK)
text(s, 0.75, 6.5, 11.9, 0.5, [("These are simulations, not field results. The pilot replaces every assumption with a measured value.", 14, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Source: reports/impact_report.md, generated by scripts/run_impact.py. The post-harvest figure is the most that matched buyers could protect, not a measured reduction.")

# 9 ---------------------------------------------------------------- feasibility
s = base(9, "Feasible and affordable: software only, open data, a basic phone", "Feasibility and affordability")
card(s, 0.55, 1.6, 4.0, 3.4, "What it needs", [
    "A basic phone and mobile data",
    "No hardware: soil moisture is modelled",
    "Free weather, soil and price data",
    "Phone + PIN login (hashed), no public farmer list"], GREEN, 14)
card(s, 4.67, 1.6, 4.0, 3.4, "What it costs", [
    "Language model: about US$0.05 per farmer per month (30 chat turns, ~1,500 tokens each; token count is an estimate)",
    "Translation and audio: free endpoints in the pilot; Bhashini (government, free) is the scale path"], AMBER, 14)
card(s, 8.79, 1.6, 4.0, 3.4, "Who pays (to test)", [
    "Farmers: free",
    "Buyers and processors: demand matching and supply-coming view",
    "Programmes (Atal Bhujal, PM-KUSUM agencies): village water and pump analytics"], BLUE, 14)
rect(s, 0.55, 5.2, 12.24, 1.7, WHITE, GREY)
text(s, 0.75, 5.25, 11.9, 1.65, [("Pilot design", 15, True, DARK)] + bullets([
    "60–80 smallholders in one Yavatmal cluster through an FPO, with the local KVK, 3–5 vendors and the gram panchayat",
    "Matched groups (with and without the schedule), pump hour or flow meters, harvest weight and price received",
    "Not yet costed: hosting, speech-to-text rate, and the field agent, the main real pilot cost"], 13, INK, 4))
notes(s, "Be upfront about what is not costed. Unit economics are hypotheses until the pilot.")

# 10 --------------------------------------------------------------- sustainability
s = base(10, "Sustainability: less pumping, less groundwater drawn, crops matched to the water that exists", "Sustainability")
card(s, 0.55, 1.6, 4.0, 3.2, "Energy", [
    "Fewer pump hours, so less electricity or diesel per unit of crop",
    "Pump sized to the real peak need: 3 HP drip vs 5 HP flood for 2 acres of cotton or wheat"], AMBER, 14)
card(s, 4.67, 1.6, 4.0, 3.2, "Water", [
    "Schedule and efficient method replace calendar flooding",
    "Village roll-up shows demand against supply before the season, where groundwater is protected"], TEAL, 14)
card(s, 8.79, 1.6, 4.0, 3.2, "Emissions", [
    "Energy saved × CEA grid factor (0.727 kg CO₂/kWh)",
    "Cotton, per acre per season: 77 to 308 kg CO₂ across lift and efficiency assumptions"], GREEN, 14)
tbl = s.shapes.add_table(4, 4, Inches(0.55), Inches(5.0), Inches(7.6), Inches(1.9)).table
data = [("Cotton, drip + schedule (kWh/acre)", "Lift 30 m", "Lift 40 m", "Lift 60 m"),
        ("Pump 25%", "212", "282", "423"), ("Pump 35%", "151", "202", "302"), ("Pump 50%", "106", "141", "212")]
for r, row in enumerate(data):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(12)
        run.font.name = "Calibri"
        run.font.bold = (r == 0 or c == 0)
        run.font.color.rgb = WHITE if r == 0 else INK
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK if r == 0 else (LIGHT if r % 2 else WHITE)
tbl.columns[0].width = Inches(3.1)
for c in (1, 2, 3):
    tbl.columns[c].width = Inches(1.5)
text(s, 8.4, 5.0, 4.4, 1.9, [("Energy scales with lift and inversely with pump efficiency, so the farmer's own pump details replace these assumptions.", 13, False, INK, 8),
                             ("Not modelled: waterlogging, heat stress, pests.", 12, False, GREY)])
notes(s, "Source: reports/impact_report.md sections 3 and 7.")

# 11 --------------------------------------------------------------- roadmap
s = base(11, "Roadmap: ideation now, working prototype next, a measured pilot after", "Implementation roadmap and future scope")
phases = [("Ideation (now)", GREEN, ["Working demo: GPS location, soil, 30-season crop ranking, My farm schedule, Sell, voice chat", "30 crops, Maharashtra rows", "Cited impact simulation"]),
          ("Prototype round (next)", TEAL, ["State-specific crop rows beyond Maharashtra, from official data", "Native-speaker review of Marathi, Hindi, Telugu", "Test on low-end phones; Bhashini voice for all 22 languages"]),
          ("Pilot (season 1)", AMBER, ["60–80 farmers, one FPO, KVK, 3–5 vendors", "Meters on pumps; matched groups", "Replace every assumption with a measurement"]),
          ("Scale", BLUE, ["Other Maharashtra districts, then Atal Bhujal states", "Open sign-up, state by state", "Panchayat and FPO dashboards"])]
for i, (h, c, b) in enumerate(phases):
    x = 0.55 + i * 3.1
    rect(s, x, 1.6, 2.95, 0.55, c)
    text(s, x, 1.63, 2.95, 0.5, [(h, 16, True, WHITE)], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    rect(s, x, 2.15, 2.95, 2.55, LIGHT)
    text(s, x + 0.1, 2.22, 2.75, 2.45, bullets(b, 14, INK, 6))
rect(s, 0.55, 4.9, 12.24, 2.0, WHITE, GREEN)
text(s, 0.75, 4.95, 11.9, 1.95, [("Future scope (ideas, not built)", 16, True, DARK)] + bullets([
    "Soil-moisture probe and pump relay that follow the schedule (design only, advice mode first, manual override always wins)",
    "SMS or phone-call (IVR) access for farmers without smartphones; offline-friendly pages",
    "Vendor payment record and, later, escrow; vendor reputation from pilot data",
    "A labelled machine-learning second opinion beside the simulation; seasonal-forecast inputs once skill is proven locally"], 13, INK, 4))
notes(s, "Ideation round: show direction. The prototype round delivers items in the second column.")

# 12 --------------------------------------------------------------- team and limits
s = base(12, "Team, and what we state plainly", "Team introduction")
for i in range(4):
    x = 0.55 + i * 3.1
    rect(s, x, 1.65, 2.95, 2.2, LIGHT)
    rect(s, x + 1.0, 1.8, 0.95, 0.95, GREEN, shape=MSO_SHAPE.OVAL)
    text(s, x, 2.85, 2.95, 0.5, [("[Name]", 17, True, DARK)], align=PP_ALIGN.CENTER)
    text(s, x, 3.3, 2.95, 0.5, [("[Role and strength]", 13, False, GREY)], align=PP_ALIGN.CENTER)
rect(s, 0.55, 4.1, 12.24, 2.8, WHITE, AMBER)
text(s, 0.75, 4.15, 11.9, 2.75, [("Limits we state plainly", 16, True, DARK)] + bullets([
    "Benefit figures are simulations until the pilot measures them",
    "Only Maharashtra has state-specific rows; vegetable costs and prices are estimates and flagged",
    "The water model covers drought, not waterlogging or pests; soil NPK cannot be read from a map",
    "Hindi and Marathi text is machine-translated and needs native review; GPS and low-end phones still to be tested in the field",
    "Demo buyers and prices are synthetic"], 14, INK, 4))
notes(s, "Edit the team names, roles and add photos here before submitting.")

prs.save(OUT)
print("saved", OUT)
